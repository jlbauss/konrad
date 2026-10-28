# SPDX-FileCopyrightText: 2026 Jan-Luca Bauß
# SPDX-License-Identifier: AGPL-3.0-or-later
"""Grade pptx-create-deck: a five-slide deck from regions.csv and accounts.csv (see prompt.md)."""

import csv
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "lib"))
from abgrade import close, norm, pdf_text, run, to_pdf  # noqa: E402

from pptx import Presentation  # noqa: E402
from pptx.enum.chart import XL_CHART_TYPE  # noqa: E402

TASK = Path(__file__).resolve().parent
NAMES = [
    "quarterly-review.pptx opens", "five slides", "16:9", "titles in title placeholders, in order",
    "subtitle", "highlights bullets", "next steps bullets", "native clustered bar chart", "chart data",
    "accounts table", "umlauts intact", "speaker notes on slides 2-5", "everything on the slide",
    "LibreOffice renders five pages", "inputs untouched",
]
TITLES = ["Quarterly Review Q3 2026", "Highlights", "Revenue by region", "Top accounts", "Next steps"]
HIGHLIGHTS = ["Revenue up 9.8% on Q2", "East and West grew fastest", "South dipped after the price change"]
STEPS = ["Win back South", "Expand the West team", "Review pricing in November"]
CLUSTERED = {XL_CHART_TYPE.COLUMN_CLUSTERED, XL_CHART_TYPE.BAR_CLUSTERED}


def rows(name):
    return list(csv.reader((TASK / "input" / name).read_text(encoding="utf-8").splitlines()))


def lines(slide):
    """Non-empty paragraphs of every text frame on the slide except its title, a typed-in bullet dropped."""
    title = slide.shapes.title
    return [norm(p.text.lstrip("•▪–-* \t")) for s in slide.shapes if s.has_text_frame and (title is None or s.shape_id != title.shape_id)
            for p in s.text_frame.paragraphs if p.text.strip()]


def grade(ws_dir, c):
    path = ws_dir / "quarterly-review.pptx"
    if not c.add(NAMES[0], path.exists(), "missing" if not path.exists() else ""):
        raise RuntimeError("no quarterly-review.pptx")
    deck = Presentation(path)
    slides = list(deck.slides)
    c.add(NAMES[1], len(slides) == 5, len(slides))
    w, h = deck.slide_width, deck.slide_height
    c.add(NAMES[2], abs(w / h - 16 / 9) < 0.01, f"{w} x {h}")
    titles = [norm(s.shapes.title.text) if s.shapes.title is not None else None for s in slides]
    c.add(NAMES[3], titles == [norm(t) for t in TITLES], titles)
    by_title = dict(zip(titles, slides))

    first = slides[0] if slides else None
    c.add(NAMES[4], first is not None and "Sales team" in lines(first), lines(first) if first else None)
    hl, ns = by_title.get("Highlights"), by_title.get("Next steps")
    c.add(NAMES[5], hl is not None and lines(hl) == [norm(t) for t in HIGHLIGHTS], lines(hl) if hl else None)
    c.add(NAMES[6], ns is not None and lines(ns) == [norm(t) for t in STEPS], lines(ns) if ns else None)

    rev = by_title.get("Revenue by region")
    charts = [s.chart for s in rev.shapes if s.has_chart] if rev is not None else []
    chart = charts[0] if charts else None
    c.add(NAMES[7], chart is not None and chart.chart_type in CLUSTERED,
          chart.chart_type if chart is not None else "no chart")
    want = rows("regions.csv")
    if chart is not None:
        plot = chart.plots[0]
        series = {norm(s.name): list(s.values) for s in plot.series}
        ok = list(plot.categories) == [r[0] for r in want[1:]] and all(
            q in series and all(close(a, float(r[i])) for a, r in zip(series[q], want[1:]))
            for i, q in ((1, "Q2"), (2, "Q3")))
        c.add(NAMES[8], ok, f"{list(plot.categories)} {series}")
    else:
        c.add(NAMES[8], False, "no chart")

    top = by_title.get("Top accounts")
    tables = [s.table for s in top.shapes if s.has_table] if top is not None else []
    got = [[cell.text.strip() for cell in row.cells] for row in tables[0].rows] if tables else []
    acc = rows("accounts.csv")
    c.add(NAMES[9], len(got) == len(acc) and got[0] == acc[0] and all(
        g[:2] == a[:2] and close(g[2].replace(",", "."), a[2]) for g, a in zip(got[1:], acc[1:])),
        got[:2])
    flat = "\n".join(cell for row in got for cell in row)
    c.add(NAMES[10], all(n in flat for n in ("Müller & Söhne KG", "Café Lumière SARL", "Østerberg Trading")))

    notes = [s.notes_slide.notes_text_frame.text.strip() if s.has_notes_slide else "" for s in slides[1:5]]
    c.add(NAMES[11], len(notes) == 4 and all(notes), [n[:40] for n in notes])
    out = [(i + 1, s.name) for i, sl in enumerate(slides) for s in sl.shapes
           if s.left is not None and (s.left < 0 or s.top < 0 or s.left + s.width > w * 1.01
                                      or s.top + s.height > h * 1.01)]
    c.add(NAMES[12], not out, out)
    pdf = to_pdf(path)
    c.add(NAMES[13], pdf is not None and pdf_text(pdf).count("\f") == 5, "no PDF" if pdf is None else "")
    c.add(NAMES[14], all((ws_dir / n).read_bytes() == (TASK / "input" / n).read_bytes()
                         for n in ("regions.csv", "accounts.csv")))


run(grade, NAMES)
