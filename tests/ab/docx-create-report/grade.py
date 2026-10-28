# SPDX-FileCopyrightText: 2026 Jan-Luca Bauß
# SPDX-License-Identifier: AGPL-3.0-or-later
"""Grade docx-create-report: a status report built from milestones.csv (see prompt.md)."""

import csv
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "lib"))
from abgrade import list_kind, norm, pdf_text, run, to_pdf  # noqa: E402

from docx import Document  # noqa: E402
from docx.shared import Mm  # noqa: E402

TASK = Path(__file__).resolve().parent
NAMES = [
    "status-report.docx opens", "title first", "four headings, in order", "summary paragraph",
    "table header", "table rows", "table header bold", "risks are a bulleted list",
    "next steps are a numbered list", "page X of Y fields in the footer", "A4", "umlauts intact",
    "LibreOffice renders the page numbers", "milestones.csv untouched",
]
SECTIONS = ["Summary", "Milestones", "Risks", "Next steps"]
RISKS = ["Load test results may force a database upgrade", "Two trainers are on leave in November",
         "Branch rollout depends on the new VPN"]
STEPS = ["Rerun the load test", "Decide on the database upgrade", "Confirm the training dates",
         "Freeze the go-live checklist"]


def same_date(text, iso):
    y, m, d = iso.split("-")
    return re.sub(r"\D", "", text) in (y + m + d, d + m + y, m + d + y)


def header_bold(table):
    """Every header run is bold, directly or by its style, or the table style bolds the first row."""
    runs = [(p, r) for cell in table.rows[0].cells for p in cell.paragraphs for r in p.runs if r.text.strip()]
    if runs and all(r.bold or (r.style is not None and r.style.font.bold) or (p.style is not None and p.style.font.bold)
                    for p, r in runs):
        return True
    style = table.style
    return style is not None and bool(style.element.xpath('.//w:tblStylePr[@w:type="firstRow"]//w:b'))


def grade(ws_dir, c):
    path = ws_dir / "status-report.docx"
    if not c.add(NAMES[0], path.exists(), "missing" if not path.exists() else ""):
        raise RuntimeError("no status-report.docx")
    doc = Document(path)
    paras = [p for p in doc.paragraphs if p.text.strip()]
    first = paras[0]
    c.add(NAMES[1], norm(first.text) == norm("Project Status – Q3 2026")
          and first.style.name in ("Title", "Heading 1"), f"{first.text!r} ({first.style.name})")

    heads = [p for p in paras[1:] if p.style.name.startswith("Heading")]
    c.add(NAMES[2], [norm(p.text) for p in heads] == [norm(s) for s in SECTIONS],
          [(p.text, p.style.name) for p in heads])
    at = {norm(p.text): i for i, p in enumerate(paras) if p in heads}

    def section(name, nxt):
        a, b = at.get(norm(name)), at.get(norm(nxt)) if nxt else len(paras)
        return paras[a + 1:b] if a is not None and b is not None else []

    summary = section("Summary", "Milestones")
    c.add(NAMES[3], any("load test" in p.text.lower() for p in summary), [p.text[:80] for p in summary])

    want = list(csv.reader((TASK / "input/milestones.csv").read_text(encoding="utf-8").splitlines()))
    table = doc.tables[0] if doc.tables else None
    if table is None:
        raise RuntimeError("no table")
    rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
    c.add(NAMES[4], rows[0] == want[0], rows[0])
    body_ok = len(rows) - 1 == len(want) - 1 and all(
        r[:2] == w[:2] and same_date(r[2], w[2]) and r[3] == w[3] for r, w in zip(rows[1:], want[1:]))
    c.add(NAMES[5], body_ok, f"{len(rows) - 1} rows; first {rows[1] if len(rows) > 1 else None}")
    c.add(NAMES[6], header_bold(table))

    risks = section("Risks", "Next steps")
    c.add(NAMES[7], [norm(p.text) for p in risks] == [norm(t) for t in RISKS]
          and all(list_kind(doc, p) == "bullet" for p in risks),
          [(p.text[:40], list_kind(doc, p)) for p in risks])
    steps = section("Next steps", None)
    c.add(NAMES[8], [norm(p.text) for p in steps] == [norm(t) for t in STEPS]
          and all(list_kind(doc, p) == "number" for p in steps),
          [(p.text[:40], list_kind(doc, p)) for p in steps])

    footers = "".join(s.footer._element.xml for s in doc.sections)
    c.add(NAMES[9], re.search(r"\bPAGE\b", footers) and "NUMPAGES" in footers)
    sec = doc.sections[0]
    c.add(NAMES[10], abs(sec.page_width - Mm(210)) < Mm(2) and abs(sec.page_height - Mm(297)) < Mm(2),
          f"{sec.page_width.mm:.0f} x {sec.page_height.mm:.0f} mm")
    text = "\n".join(cell.text for row in table.rows for cell in row.cells)
    c.add(NAMES[11], "Jürgen Weiß" in text and "Özlem Kaya" in text and "Marie Groß" in text)
    pdf = to_pdf(path)
    c.add(NAMES[12], pdf is not None and re.search(r"Page 1 of \d", pdf_text(pdf)),
          "no PDF" if pdf is None else "")
    c.add(NAMES[13], (ws_dir / "milestones.csv").read_bytes() == (TASK / "input/milestones.csv").read_bytes())


run(grade, NAMES)
