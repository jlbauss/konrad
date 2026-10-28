# SPDX-FileCopyrightText: 2026 Jan-Luca Bauß
# SPDX-License-Identifier: AGPL-3.0-or-later
"""Grade pptx-edit-deck: a rename, a chart update, a deleted and a moved slide in deck.pptx (see prompt.md)."""

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "lib"))
from abgrade import close, norm, pdf_text, run, to_pdf  # noqa: E402

from pptx import Presentation  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE_TYPE  # noqa: E402

TASK = Path(__file__).resolve().parent
NAMES = [
    "deck-v2.pptx opens", "slides in the new order", "old name gone", "new name in title, chart and notes",
    "chart still a native chart", "chart categories Q1-Q4", "chart values", "team photo kept", "team notes kept",
    "everything else unchanged", "LibreOffice renders five pages", "deck.pptx untouched",
]
ORDER = ["Aurora Product Update", "Agenda", "Aurora adoption", "Team", "Summary"]


def texts(shapes):
    for s in shapes:
        if s.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from texts(s.shapes)
        if s.has_text_frame:
            yield from (p.text for p in s.text_frame.paragraphs if p.text.strip())
        if getattr(s, "has_table", False) and s.has_table:
            yield from (cell.text for row in s.table.rows for cell in row.cells if cell.text.strip())


def notes(slide):
    return slide.notes_slide.notes_text_frame.text.strip() if slide.has_notes_slide else ""


def chart_of(slide):
    charts = [s.chart for s in slide.shapes if s.has_chart]
    return charts[0] if charts else None


def chart_text(chart):
    """Everything a viewer can read off the chart: its title and series names, from the chart part's XML."""
    return chart.part.blob.decode("utf-8", "replace") if chart is not None else ""


def grade(ws_dir, c):
    src = TASK / "input/deck.pptx"
    path = ws_dir / "deck-v2.pptx"
    if not c.add(NAMES[0], path.exists(), "missing" if not path.exists() else ""):
        raise RuntimeError("no deck-v2.pptx")
    deck, orig = Presentation(path), Presentation(src)
    slides = list(deck.slides)
    titles = [norm(s.shapes.title.text) if s.shapes.title is not None else None for s in slides]
    c.add(NAMES[1], titles == ORDER, titles)
    by_title = dict(zip(titles, slides))

    every = "\n".join(t for s in slides for t in [*texts(s.shapes), notes(s), chart_text(chart_of(s))])
    c.add(NAMES[2], "Alpha" not in every, every.count("Alpha"))
    adoption = by_title.get("Aurora adoption")
    chart = chart_of(adoption) if adoption is not None else None
    places = [titles[0] if titles else "", chart_text(chart),
              *(notes(by_title[t]) if t in by_title else "" for t in ("Agenda", "Aurora adoption", "Team"))]
    c.add(NAMES[3], all("Aurora" in p for p in places), [("Aurora" in p) for p in places])

    c.add(NAMES[4], chart is not None, "" if chart is not None else "no chart on the adoption slide")
    plot = chart.plots[0] if chart is not None else None
    c.add(NAMES[5], plot is not None and list(plot.categories) == ["Q1", "Q2", "Q3", "Q4"],
          list(plot.categories) if plot is not None else None)
    values = list(plot.series[0].values) if plot is not None and len(plot.series) else []
    c.add(NAMES[6], len(values) == 4 and all(close(a, b) for a, b in zip(values, [820, 1210, 1640, 2100])), values)

    team = by_title.get("Team")
    c.add(NAMES[7], team is not None and any(s.shape_type == MSO_SHAPE_TYPE.PICTURE for s in team.shapes))
    c.add(NAMES[8], team is not None and norm(notes(team)) == "Introduce the Aurora team leads",
          notes(team) if team is not None else None)

    def content(s, renamed):
        lines = [norm(t.replace("Alpha", "Aurora") if renamed else t) for t in texts(s.shapes)]
        return lines, norm(notes(s).replace("Alpha", "Aurora") if renamed else notes(s))

    before = {content(s, True)[0][0]: content(s, True) for s in orig.slides}
    after = {content(s, False)[0][0]: content(s, False) for s in slides if list(texts(s.shapes))}
    changed = [t for t in after if before.get(t) != after[t]]
    c.add(NAMES[9], set(after) == set(before) - {"Old roadmap"} and not changed, changed or sorted(after))
    pdf = to_pdf(path)
    c.add(NAMES[10], pdf is not None and pdf_text(pdf).count("\f") == 5, "no PDF" if pdf is None else "")
    c.add(NAMES[11], (ws_dir / "deck.pptx").read_bytes() == src.read_bytes())


run(grade, NAMES)
