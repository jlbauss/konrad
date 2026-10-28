# SPDX-FileCopyrightText: 2026 Jan-Luca Bauß
# SPDX-License-Identifier: AGPL-3.0-or-later
"""Regenerate input/deck.pptx (committed, so every trial starts from the same bytes).

Run in a konrad image. The product name hides in a title, bullets, a text box,
a chart title and three speaker notes; the chart is native, so its data lives
in an embedded workbook that has to move with any edit.

    podman run --rm -v "$PWD/tests/ab/pptx-edit-deck:/t:z" --entrypoint /opt/venv/bin/python \\
        konrad:local /t/make_input.py /t/input/deck.pptx
"""

import io
import sys

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches, Pt


def bullets(deck, title, items, notes=None):
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = title
    body = slide.placeholders[1].text_frame
    body.text = items[0]
    for item in items[1:]:
        body.add_paragraph().text = item
    if notes:
        slide.notes_slide.notes_text_frame.text = notes
    return slide


def photo():
    img = Image.new("RGB", (400, 300), "#d9e2ec")
    ImageDraw.Draw(img).ellipse((140, 60, 260, 180), fill="#486581")
    buf = io.BytesIO()
    img.save(buf, "PNG")
    buf.seek(0)
    return buf


def main(out):
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)

    title = deck.slides.add_slide(deck.slide_layouts[0])
    title.shapes.title.text = "Alpha Product Update"
    title.placeholders[1].text = "October 2026"

    bullets(deck, "Agenda", ["Where Alpha stands", "Adoption", "Team"], "Keep the Alpha intro short.")

    adoption = deck.slides.add_slide(deck.slide_layouts[5])
    adoption.shapes.title.text = "Alpha adoption"
    data = CategoryChartData()
    data.categories = ["Q1", "Q2", "Q3"]
    data.add_series("Active users", (820, 1210, 1500))
    chart = adoption.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1.5), Inches(1.6),
                                      Inches(10), Inches(5.2), data).chart
    chart.has_title = True
    chart.chart_title.text_frame.text = "Alpha active users"
    chart.chart_title.text_frame.paragraphs[0].runs[0].font.size = Pt(18)
    adoption.notes_slide.notes_text_frame.text = "Adoption grew every quarter since Alpha launched."

    bullets(deck, "Old roadmap", ["Beta in Q1", "Mobile app in Q2"])
    bullets(deck, "Summary", ["Alpha is on track", "Hiring two engineers"])

    team = deck.slides.add_slide(deck.slide_layouts[5])
    team.shapes.title.text = "Team"
    team.shapes.add_picture(photo(), Inches(1), Inches(1.8), width=Inches(4))
    box = team.shapes.add_textbox(Inches(6), Inches(2.5), Inches(6), Inches(1)).text_frame
    box.text = "Ask the Alpha team: team@example.com"
    team.notes_slide.notes_text_frame.text = "Introduce the Alpha team leads."
    deck.save(out)


main(sys.argv[1])
