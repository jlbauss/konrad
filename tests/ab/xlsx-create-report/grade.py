# SPDX-FileCopyrightText: 2026 Jan-Luca Bauß
# SPDX-License-Identifier: AGPL-3.0-or-later
"""Grade xlsx-create-report: report.xlsx built from sales.csv (see prompt.md)."""

import csv
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "lib"))
from abgrade import close, is_formula, recalc, run  # noqa: E402

from openpyxl import load_workbook  # noqa: E402

TASK = Path(__file__).resolve().parent
NAMES = [
    "report.xlsx opens", "sheets Data and Summary", "Data has all 36 rows",
    "Revenue cells are formulas", "Revenue values correct", "Summary lists each region once",
    "region totals are formulas", "region totals correct", "Total row correct",
    "currency format, two decimals", "header rows bold", "Data header frozen",
    "non-ASCII text intact", "sales.csv untouched",
]


def money(fmt):
    return "0.00" in fmt and (any(s in fmt for s in "$€£¥") or "[$" in fmt)


def grade(ws_dir, c):
    rows = list(csv.DictReader(open(TASK / "input/sales.csv", encoding="utf-8")))
    revenue = [int(r["Units"]) * float(r["Unit Price"]) for r in rows]
    by_region = {}
    for r, v in zip(rows, revenue):
        by_region[r["Region"]] = by_region.get(r["Region"], 0) + v

    path = ws_dir / "report.xlsx"
    if not c.add(NAMES[0], path.exists(), "missing" if not path.exists() else ""):
        raise RuntimeError("no report.xlsx")
    wb = load_workbook(path)
    vb = load_workbook(recalc(path), data_only=True)
    if not c.add(NAMES[1], {"Data", "Summary"} <= set(wb.sheetnames), wb.sheetnames):
        raise RuntimeError("sheets missing")
    data, vdata, summ, vsumm = wb["Data"], vb["Data"], wb["Summary"], vb["Summary"]

    header = [str(x.value or "").strip().lower() for x in data[1]]
    rcol = next((i + 1 for i, h in enumerate(header) if "revenue" in h), None)
    body = [r for r in range(2, data.max_row + 1) if data.cell(r, 1).value not in (None, "")]
    c.add(NAMES[2], len(body) == 36, f"{len(body)} rows")
    if rcol is None:
        raise RuntimeError(f"no Revenue column in {header}")
    cells = [data.cell(r, rcol) for r in body]
    c.add(NAMES[3], all(is_formula(x.value) for x in cells),
          f"{sum(is_formula(x.value) for x in cells)}/{len(cells)} formulas")
    got = sorted(float(vdata.cell(r, rcol).value or 0) for r in body)
    c.add(NAMES[4], len(got) == 36 and all(close(a, b) for a, b in zip(got, sorted(revenue))),
          f"sum {sum(got):.2f} vs {sum(revenue):.2f}")

    region_rows = {}
    for r in range(1, summ.max_row + 1):
        label = str(summ.cell(r, 1).value or "").strip()
        if label in by_region:
            region_rows.setdefault(label, []).append(r)
    c.add(NAMES[5], sorted(region_rows) == sorted(by_region) and all(len(v) == 1 for v in region_rows.values()),
          {k: len(v) for k, v in region_rows.items()})
    rr = [v[0] for v in region_rows.values()]
    c.add(NAMES[6], rr and all(is_formula(summ.cell(r, 2).value) for r in rr),
          [summ.cell(r, 2).value for r in rr][:2])
    c.add(NAMES[7], rr and all(close(vsumm.cell(region_rows[k][0], 2).value, v) for k, v in by_region.items() if k in region_rows),
          {k: vsumm.cell(region_rows[k][0], 2).value for k in region_rows})
    total = next((r for r in range(1, summ.max_row + 1)
                  if str(summ.cell(r, 1).value or "").strip().lower() in ("total", "grand total")), None)
    c.add(NAMES[8], total is not None and close(vsumm.cell(total, 2).value, sum(revenue)),
          f"{vsumm.cell(total, 2).value if total else 'no Total row'} vs {sum(revenue):.2f}")

    fmts = {x.number_format for x in cells} | {summ.cell(r, 2).number_format for r in rr + ([total] if total else [])}
    c.add(NAMES[9], fmts and all(money(f) for f in fmts), sorted(fmts))
    first = [x for x in summ[1] if x.value not in (None, "")]
    summ_has_header = str(summ.cell(1, 1).value or "").strip() not in by_region
    bold = all(x.font.bold for x in data[1] if x.value not in (None, "")) and \
        (not summ_has_header or all(x.font.bold for x in first))
    c.add(NAMES[10], bold, f"Summary header row: {summ_has_header}")
    c.add(NAMES[11], bool(re.fullmatch(r"[A-Z]+2", data.freeze_panes or "")), data.freeze_panes)
    texts = {str(x.value) for row in data.iter_rows() for x in row}
    c.add(NAMES[12], "Zubehör" in texts, "found" if "Zubehör" in texts else "Zubehör missing")
    c.add(NAMES[13], (ws_dir / "sales.csv").read_bytes() == (TASK / "input/sales.csv").read_bytes())


run(grade, NAMES)
