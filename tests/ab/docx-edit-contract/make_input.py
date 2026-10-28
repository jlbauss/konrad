# SPDX-FileCopyrightText: 2026 Jan-Luca Bauß
# SPDX-License-Identifier: AGPL-3.0-or-later
"""Regenerate input/contract.docx (committed, so every trial starts from the same bytes).

Run in a konrad image. The client's name appears six times: header, footer,
a bold run, a run split in two the way Word's editing leaves text, a table
cell and the signature line. The logo is the part a careless re-save drops.

    podman run --rm -v "$PWD/tests/ab/docx-edit-contract:/t:z" --entrypoint /opt/venv/bin/python \\
        konrad:local /t/make_input.py /t/input/contract.docx
"""

import io
import sys

from docx import Document
from docx.shared import Mm
from PIL import Image, ImageDraw

NAME = "Nordlicht GmbH"


def logo():
    img = Image.new("RGB", (240, 80), "#1f4e79")
    ImageDraw.Draw(img).text((20, 30), "NORDLICHT", fill="white")
    buf = io.BytesIO()
    img.save(buf, "PNG")
    buf.seek(0)
    return buf


def main(out):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.header.paragraphs[0].text = f"{NAME} – Service Agreement"
    sec.footer.paragraphs[0].text = f"{NAME} · Confidential"

    doc.add_picture(logo(), width=Mm(40))
    doc.add_paragraph("Service Agreement", style="Title")
    p = doc.add_paragraph("between ")
    p.add_run(NAME).bold = True
    p.add_run(', Hafenstraße 12, 20457 Hamburg (the "Client"), and Brightline Consulting Ltd (the "Contractor").')

    doc.add_heading("1. Scope", level=1)
    p = doc.add_paragraph("The Contractor supports ")
    p.add_run("Nordl")
    p.add_run("icht GmbH")
    p.add_run(" in migrating its billing system to the new platform.")

    doc.add_heading("2. Term", level=1)
    doc.add_paragraph("This agreement runs from 1 January 2027 for twelve months and renews by one year unless "
                      "either party gives notice three months before it ends.")

    doc.add_heading("3. Fees", level=1)
    table = doc.add_table(rows=4, cols=3, style="Table Grid")
    for r, row in enumerate([["Item", "Rate", "Unit"], ["Consulting", "950 EUR", "day"],
                             ["Travel", "at cost", "trip"], ["Billed to", NAME, "monthly"]]):
        for col, text in enumerate(row):
            table.cell(r, col).text = text

    doc.add_heading("4. Payment", level=1)
    doc.add_paragraph("Invoices are payable within 30 days of receipt.")

    doc.add_heading("5. Liability", level=1)
    doc.add_paragraph("Liability is capped at the total fees paid in the preceding twelve months.")

    doc.add_heading("6. Signatures", level=1)
    doc.add_paragraph(f"For {NAME}: ______________________")
    doc.add_paragraph("For Brightline Consulting Ltd: ______________________")
    doc.core_properties.author = "Brightline Consulting Ltd"
    doc.save(out)


main(sys.argv[1])
