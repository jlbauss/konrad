# SPDX-FileCopyrightText: 2026 Jan-Luca Bauß
# SPDX-License-Identifier: AGPL-3.0-or-later
"""Regenerate input/review.docx (committed, so every trial starts from the same bytes).

Run in a konrad image. Two reviewers' tracked changes sit in the body and in a
table, some side by side in one paragraph, plus three comments. The expected
result of accepting one reviewer and rejecting the other is in grade.py.

    podman run --rm -v "$PWD/tests/ab/docx-review-revisions:/t:z" --entrypoint /opt/venv/bin/python \\
        konrad:local /t/make_input.py /t/input/review.docx
"""

import sys

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ANNA, BEN = "Anna Berg", "Ben Ott"
DATE = {ANNA: "2026-10-20T09:12:00Z", BEN: "2026-10-21T14:40:00Z"}
ids = iter(range(100, 200))

# Each paragraph is a list of (kind, author, text); kind is "", "ins" or "del".
BODY = [
    [("", None, "The pilot ran in "), ("del", ANNA, "four"), ("ins", ANNA, "five"),
     ("", None, " branches from June to September.")],
    [("", None, "Staff rated the new system "), ("ins", BEN, "very "), ("", None, "positively"),
     ("del", BEN, ", with few complaints"), ("", None, ".")],
    [("", None, "Support tickets fell by "), ("del", ANNA, "20"), ("ins", ANNA, "23"), ("", None, " percent"),
     ("ins", BEN, " overall"), ("", None, ".")],
    [("", None, "The rollout to all branches "), ("del", BEN, "is planned for"), ("ins", BEN, "will slip to"),
     ("", None, " January.")],
]
NEXT = [("", None, "Next review: "), ("ins", ANNA, "15 "), ("", None, "November.")]


def run(text, deleted=False):
    r = OxmlElement("w:r")
    t = OxmlElement("w:delText" if deleted else "w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    r.append(t)
    return r


def fill(p, segments):
    """Append the segments to python-docx paragraph p as runs and tracked changes."""
    for kind, author, text in segments:
        if not kind:
            p._p.append(run(text))
            continue
        w = OxmlElement(f"w:{kind}")
        w.set(qn("w:id"), str(next(ids)))
        w.set(qn("w:author"), author)
        w.set(qn("w:date"), DATE[author])
        w.append(run(text, deleted=kind == "del"))
        p._p.append(w)
    return p


def main(out):
    doc = Document()
    doc.add_paragraph("Pilot Report: New Branch System", style="Title")
    doc.add_heading("Results", level=1)
    paras = [fill(doc.add_paragraph(), segs) for segs in BODY]
    doc.add_heading("Status", level=1)
    table = doc.add_table(rows=3, cols=2, style="Table Grid")
    table.cell(0, 0).text, table.cell(0, 1).text = "Item", "Value"
    table.cell(1, 0).text = "Phase"
    fill(table.cell(1, 1).paragraphs[0], [("del", ANNA, "Pilot"), ("ins", ANNA, "Rollout")])
    table.cell(2, 0).text = "Budget owner"
    fill(table.cell(2, 1).paragraphs[0], [("", None, "Finance"), ("ins", BEN, " (TBD)")])
    nxt = fill(doc.add_paragraph(), NEXT)

    doc.add_comment(paras[0].runs[0], text="Five, not four – see the branch list.", author=ANNA, initials="AB")
    doc.add_comment(paras[3].runs[0], text="Do we have a new date?", author=BEN, initials="BO")
    doc.add_comment(nxt.runs[0], text="Confirm after the steering meeting.", author=ANNA, initials="AB")
    doc.save(out)


main(sys.argv[1])
