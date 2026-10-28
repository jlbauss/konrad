# SPDX-FileCopyrightText: 2026 Jan-Luca Bauß
# SPDX-License-Identifier: AGPL-3.0-or-later
"""Grade xlsx-csv-locale: a cp1252, semicolon, decimal-comma CSV turned into export.xlsx."""

import datetime
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "lib"))
from abgrade import close, is_formula, recalc, run  # noqa: E402

from openpyxl import load_workbook  # noqa: E402

TASK = Path(__file__).resolve().parent
NAMES = [
    "export.xlsx opens", "headers kept", "one row per CSV line", "dates are real dates",
    "amounts are correct numbers", "VAT rates are numbers", "customer names exact",
    "sum row is a formula", "sum correct", "export.csv untouched",
]


def source():
    rows = []
    lines = (TASK / "input/export.csv").read_text(encoding="cp1252").splitlines()
    for line in lines[1:]:
        date, kunde, betrag, rate = line.split(";")
        rows.append((datetime.datetime.strptime(date, "%d.%m.%Y").date(), kunde,
                     float(betrag.replace(".", "").replace(",", ".")), float(rate.rstrip("%")) / 100))
    return lines[0].split(";"), rows


def grade(ws_dir, c):
    header, src = source()
    path = ws_dir / "export.xlsx"
    if not c.add(NAMES[0], path.exists(), "missing" if not path.exists() else ""):
        raise RuntimeError("no export.xlsx")
    wb = load_workbook(path)
    ws = wb.worksheets[0]
    vs = load_workbook(recalc(path), data_only=True)[wb.sheetnames[0]]

    got_header = [str(x.value or "").strip() for x in ws[1]][:4]
    c.add(NAMES[1], got_header == header, got_header)
    # Data rows: a customer in column B and no formula anywhere (the sum row has one).
    body = [r for r in range(2, ws.max_row + 1)
            if str(ws.cell(r, 2).value or "").strip()
            and not any(is_formula(ws.cell(r, col).value) for col in range(1, 6))]
    c.add(NAMES[2], len(body) == len(src), f"{len(body)} rows, CSV has {len(src)}")
    pairs = list(zip(body, src))
    c.add(NAMES[3], len(pairs) == len(src) and all(
        isinstance(ws.cell(r, 1).value, (datetime.date, datetime.datetime)) and
        (ws.cell(r, 1).value.date() if isinstance(ws.cell(r, 1).value, datetime.datetime) else ws.cell(r, 1).value) == s[0]
        for r, s in pairs), [ws.cell(r, 1).value for r in body[:2]])
    c.add(NAMES[4], len(pairs) == len(src) and all(
        isinstance(ws.cell(r, 3).value, (int, float)) and close(ws.cell(r, 3).value, s[2]) for r, s in pairs),
        [ws.cell(r, 3).value for r in body[:3]])
    c.add(NAMES[5], len(pairs) == len(src) and all(
        isinstance(ws.cell(r, 4).value, (int, float)) and
        (close(ws.cell(r, 4).value, s[3], 1e-6) or close(ws.cell(r, 4).value, s[3] * 100, 1e-6)) for r, s in pairs),
        [ws.cell(r, 4).value for r in body[:3]])
    c.add(NAMES[6], [ws.cell(r, 2).value for r in body] == [s[1] for s in src][:len(body)] and len(body) == len(src),
          sorted({str(ws.cell(r, 2).value) for r in body})[:5])

    last = max(body) if body else 1
    sums = [(r, col) for r in range(last + 1, ws.max_row + 1) for col in range(1, 6)
            if is_formula(ws.cell(r, col).value) and "SUM" in ws.cell(r, col).value.upper()]
    c.add(NAMES[7], bool(sums), sums[:2])
    total = sum(s[2] for s in src)
    c.add(NAMES[8], any(close(vs.cell(r, col).value, total) for r, col in sums),
          f"{[vs.cell(r, col).value for r, col in sums]} vs {total:.2f}")
    c.add(NAMES[9], (ws_dir / "export.csv").read_bytes() == (TASK / "input/export.csv").read_bytes())


run(grade, NAMES)
