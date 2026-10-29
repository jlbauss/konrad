# SPDX-FileCopyrightText: 2026 Jan-Luca Bauß
# SPDX-License-Identifier: AGPL-3.0-or-later
"""Blind judging for the A/B harness; scripts/ab-judge.sh runs it in a konrad image.

blind:  each pair of runs (a-n, b-n) of a results dir becomes two packets,
        X/Y swapped, under random ids: the task, its input, and per side
        every deliverable as page renders, text and speaker notes, never the
        file itself or the agent's notes, which could say how it was made.
        The key goes into the results dir, the packets and the batches (one
        task and one order per batch, so no judge sees both halves of a
        pair) into a judge dir beside it.
report: unblinds the verdict.json files the judges wrote, combines both
        orders (a pair counts for a side only when both agree) and writes
        report/README.md with img/ beside it: a short study with the graded
        checks, the verdicts and the clearest pairs side by side, for the
        MR to commit under tests/ab/reports/.
"""

import argparse
import datetime
import email
import email.policy
import html
import json
import random
import re
import secrets
import shutil
import subprocess
import tempfile
from collections import defaultdict
from pathlib import Path

from abdump import outline
from PIL import Image

DOCS = {".docx", ".pptx", ".xlsx", ".pdf"}
TEXTS = {".txt", ".csv", ".md", ".json", ".py", ".html"}
IMAGES = {".png", ".jpg", ".jpeg"}
MAIL = {".eml"}
CRITERIA = ("request", "fidelity", "presentation")
MAX_PAGES = 12


def files(ws):
    """Every file the run left under ws, relative, without hidden ones (konrad's .agent/ among them) and caches."""
    if not ws.is_dir():
        return []
    rels = (f.relative_to(ws) for f in ws.rglob("*") if f.is_file())
    return sorted(r for r in rels if not any(p.startswith(".") or p == "__pycache__" for p in r.parts))


def deliverables(ws, inputs):
    """What the run delivered, other than unchanged inputs: documents, emails, images and text files at the top,
    and text files (an index, a list) in folders."""
    names = set()
    for rel in files(ws):
        kinds = DOCS | TEXTS | IMAGES | MAIL if len(rel.parts) == 1 else TEXTS
        if rel.suffix.lower() in kinds:
            src = inputs / rel
            if not (src.exists() and src.read_bytes() == (ws / rel).read_bytes()):
                names.add(str(rel))
    return names


def file_list(ws, inputs):
    """FILES.txt: the attempt's workspace as a list, so the judge sees names, folders and leftovers."""
    lines = []
    for rel in files(ws):
        src = inputs / rel
        note = " (input, unchanged)" if src.exists() and src.read_bytes() == (ws / rel).read_bytes() else ""
        lines.append(f"{rel}  {(ws / rel).stat().st_size} bytes{note}")
    return "\n".join(lines or ["(no files)"]) + "\n"


def mail_text(msg):
    """An email as a reader sees it: the headers that matter, the text, and the attachments by name."""
    lines = [f"{h}: {msg[h]}" for h in ("From", "To", "Cc", "Subject", "Date", "Message-ID", "In-Reply-To", "References")
             if msg[h] is not None]
    body = msg.get_body(preferencelist=("plain", "html"))
    lines += ["", body.get_content() if body else "(no text part)", ""]
    lines += [f"Attachment: {a.get_filename()} ({a.get_content_type()}, {len(a.get_payload(decode=True) or b'')} bytes)"
              for a in msg.iter_attachments()]
    return "\n".join(lines) + "\n"


def render(src, out, dpi):
    out.mkdir(parents=True)
    if not src.exists():
        (out / "MISSING.txt").write_text("No such file was delivered.\n")
        return
    suffix = src.suffix.lower()
    if suffix in TEXTS:
        shutil.copy(src, out / "text.txt")
        return
    if suffix in IMAGES:
        try:
            with Image.open(src) as im:
                im.thumbnail((1600, 1600))
                im.convert("RGB").save(out / "page-1.png")
        except OSError:
            (out / "MISSING.txt").write_text("The image can't be opened.\n")
        return
    if suffix in MAIL:
        msg = email.message_from_bytes(src.read_bytes(), policy=email.policy.default)
        (out / "text.txt").write_text(mail_text(msg), encoding="utf-8")
        for a in msg.iter_attachments():
            name = Path(a.get_filename() or "").name
            if Path(name).suffix.lower() in DOCS:
                doc = Path(tempfile.mkdtemp()) / name
                doc.write_bytes(a.get_payload(decode=True) or b"")
                pages(doc, out, dpi, "attachment-")
                break
        return
    try:
        shape = outline(src)
    except Exception as exc:  # noqa: BLE001 — a file its library can't read is itself a finding
        shape = {"error": f"could not read the file's structure: {type(exc).__name__}"}
    if shape is not None:
        (out / "outline.json").write_text(json.dumps(shape, ensure_ascii=False, indent=1, default=str), encoding="utf-8")
    pages(src, out, dpi)
    if suffix == ".pptx":
        from pptx import Presentation
        notes = [f"Slide {i}: " + (s.notes_slide.notes_text_frame.text.strip() if s.has_notes_slide else "")
                 for i, s in enumerate(Presentation(src).slides, 1)]
        (out / "notes.txt").write_text("\n".join(notes) + "\n", encoding="utf-8")


def pages(src, out, dpi, prefix=""):
    """Page images (page-N.png) and text of an office file or PDF, through LibreOffice for the former."""
    pdf = src
    if src.suffix.lower() != ".pdf":
        tmp = Path(tempfile.mkdtemp())
        subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(tmp), str(src)],
                       capture_output=True, timeout=180, check=False, env={"HOME": str(tmp), "PATH": "/usr/bin:/bin"})
        pdf = tmp / (src.stem + ".pdf")
    if not pdf.exists():
        (out / "MISSING.txt").write_text(f"LibreOffice could not open {prefix}{src.name}.\n")
        return
    subprocess.run(["pdftoppm", "-r", str(dpi), "-l", str(MAX_PAGES), "-png", str(pdf), str(out / "page")], check=True)
    subprocess.run(["pdftotext", "-layout", str(pdf), str(out / f"{prefix}text.txt")], check=True)


def runs(results, task):
    """Pair numbers with both an a- and a b- run."""
    return sorted(int(d.name[2:]) for d in (results / task).glob("a-*") if (results / task / f"b-{d.name[2:]}").is_dir())


def blind(args):
    results, judge, tasks_dir = Path(args.results), Path(args.judge), Path(args.tasks_dir)
    if (judge / "packets").exists():
        shutil.rmtree(judge / "packets")
    (judge / "packets").mkdir(parents=True)
    shutil.copy(tasks_dir / "JUDGE.md", judge / "JUDGE.md")
    key, batches = [], []
    tasks = sorted(d.name for d in results.iterdir() if d.is_dir() and (tasks_dir / d.name / "prompt.md").exists())
    if args.tasks:
        tasks = [t for t in tasks if t in args.tasks.split(",")]
    for task in tasks:
        inputs = tasks_dir / task / "input"
        by_order = {1: [], 2: []}
        for n in runs(results, task):
            names = deliverables(results / task / f"a-{n}" / "ws", inputs) | deliverables(results / task / f"b-{n}" / "ws", inputs)
            first = random.randrange(2)
            for order in (1, 2):
                x = "a" if (first + order) % 2 == 0 else "b"
                y = "b" if x == "a" else "a"
                pid = secrets.token_hex(4)
                p = judge / "packets" / pid
                p.mkdir()
                shutil.copy(tasks_dir / task / "prompt.md", p / "task.md")
                for f in sorted(inputs.iterdir()):
                    if f.suffix.lower() in DOCS | IMAGES | MAIL:
                        render(f, p / "input" / f.name, args.dpi)
                    elif f.is_dir():
                        shutil.copytree(f, p / "input" / f.name)
                    else:
                        (p / "input").mkdir(exist_ok=True)
                        shutil.copy(f, p / "input" / f.name)
                for label, side in (("X", x), ("Y", y)):
                    (p / label).mkdir()
                    (p / label / "FILES.txt").write_text(file_list(results / task / f"{side}-{n}" / "ws", inputs))
                    if not names:
                        render(Path("/nonexistent"), p / label / "deliverable", args.dpi)
                    for name in sorted(names):
                        render(results / task / f"{side}-{n}" / "ws" / name, p / label / name, args.dpi)
                key.append({"id": pid, "task": task, "pair": n, "order": order, "X": x})
                by_order[order].append(pid)
        for order in (1, 2):
            ids = by_order[order]
            random.shuffle(ids)
            for i in range(0, len(ids), args.batch):
                batches.append(" ".join(f"{args.host_judge}/packets/{pid}" for pid in ids[i:i + args.batch]))
    (results / "judge-key.json").write_text(json.dumps(key, indent=1))
    (judge / "batches.txt").write_text("\n".join(batches) + "\n")
    print(f"{len(key)} packets in {len(batches)} batches")


def verdicts(results, judge):
    """{task: {pair: {order: verdict in a/b terms}}} for every packet a judge has answered."""
    out = defaultdict(lambda: defaultdict(dict))
    for k in json.loads((results / "judge-key.json").read_text()):
        f = judge / "packets" / k["id"] / "verdict.json"
        if not f.exists():
            continue
        v = json.loads(f.read_text())
        side = {"X": k["X"], "Y": "b" if k["X"] == "a" else "a"}
        out[k["task"]][k["pair"]][k["order"]] = {
            "winner": side.get(v.get("preferred"), "tie"), "strength": v.get("strength", ""),
            "scores": {side[s]: v.get("scores", {}).get(s, {}) for s in ("X", "Y")},
            "reasons": v.get("reasons", ""), "id": k["id"], "X": k["X"]}
    return out


def summarize(pairs):
    both = {n: o for n, o in pairs.items() if 1 in o and 2 in o}
    final = {n: (o[1]["winner"] if o[1]["winner"] == o[2]["winner"] else "tie") for n, o in both.items()}
    means = {}
    for crit in CRITERIA:
        for s in "ab":
            vals = [o[k]["scores"][s].get(crit) for o in both.values() for k in (1, 2)]
            vals = [v for v in vals if isinstance(v, (int, float))]
            means[(crit, s)] = sum(vals) / len(vals) if vals else None
    return {"pairs": len(both), "consistent": sum(o[1]["winner"] == o[2]["winner"] for o in both.values()),
            "a": sum(w == "a" for w in final.values()), "b": sum(w == "b" for w in final.values()),
            "tie": sum(w == "tie" for w in final.values()), "means": means, "final": final, "both": both}


def graded(results):
    """Per task and side: the automatic checks, time, cost and tokens of the runs."""
    f = results / "results.json"
    rows = json.loads(f.read_text()) if f.exists() else []
    out = defaultdict(dict)
    for task in {r["task"] for r in rows}:
        for s in "ab":
            rs = [r for r in rows if r["task"] == task and r["side"] == s]
            if rs:
                def mean(k):
                    return sum(r.get(k) or 0 for r in rs) / len(rs)
                out[task][s] = {"runs": len(rs), "score": mean("score"),
                                "all_pass": sum(r["score"] == 1 for r in rs),
                                "cost": mean("cost"), "seconds": mean("seconds"),
                                "tokens_in": mean("tokens_in"), "tokens_out": mean("tokens_out"),
                                "skill_runs": sum(bool(set(r["skills"]) - {"quality-assurance"}) for r in rs)}
    return out


def fmt(v, spec):
    return "–" if v is None else format(v, spec)


def examples(s, limit=3):
    """The clearest pairs: consistent verdicts first, the stronger first, both directions represented."""
    rank = {"strong": 3, "clear": 2, "slight": 1}
    order = sorted(s["both"], key=lambda n: (s["final"][n] == "tie",
                                             -sum(rank.get(s["both"][n][k]["strength"], 0) for k in (1, 2))))
    picked = []
    for want in ("a", "b", "tie"):
        picked += [n for n in order if s["final"][n] == want][:1]
    picked += [n for n in order if n not in picked]
    return picked[:limit]


def side_dirs(judge, o):
    """{a: dir, b: dir} of the X/Y renders in the order-1 packet."""
    p = judge / "packets" / o[1]["id"]
    return {o[1]["X"]: p / "X", ("b" if o[1]["X"] == "a" else "a"): p / "Y"}


def pair_image(a, b, out, gap=16, edge=1):
    """One JPEG with A's page on the left and B's on the right, on gray so a white
    page shows its edges; a missing side stays gray."""
    from PIL import Image
    ims = [Image.open(x).convert("RGB") if x else None for x in (a, b)]
    w = max(i.width for i in ims if i) + 2 * edge
    h = max(i.height for i in ims if i) + 2 * edge
    canvas = Image.new("RGB", (2 * w + gap, h), "#a8a29e")
    for k, im in enumerate(ims):
        if im:
            canvas.paste(im, (k * (w + gap) + edge, edge))
    canvas.save(out, "JPEG", quality=70, optimize=True)


def ab_terms(text, x):
    """A reviewer's reasons with X and Y replaced by the versions they stood for."""
    y = "b" if x == "a" else "a"
    return re.sub(r"\b([XY])\b", lambda m: (x if m.group(1) == "X" else y).upper(), text)


def call(s):
    """How a task came out, in one phrase, and who won ('a', 'b' or '')."""
    a, b, n = s["a"], s["b"], s["pairs"] or 1
    if a == b or abs(a - b) < 2:
        return "no real difference", ""
    lead, win = ("A", "a") if a > b else ("B", "b")
    return (f"{lead} clearly better" if abs(a - b) >= n / 2 else f"{lead} somewhat better"), win


def first_sentence(text):
    m = re.match(r"(.+?[.!?])(\s|$)", text.strip().replace("\n", " "))
    return m.group(1) if m else text.strip()[:200]


def cell(text):
    """Text that can sit in a Markdown table cell or after a <summary>."""
    return html.escape(" ".join(text.split())).replace("|", "\\|")


def shows(specs):
    """--show task:run:page[,page…] → {task: (run, [pages])}: the pages an agent picked by hand."""
    out = {}
    for spec in specs:
        task, run, pages = spec.split(":")
        out[task] = (int(run), [int(x) for x in pages.split(",")])
    return out


def report(args):
    results, judge, tasks_dir = Path(args.results), Path(args.judge), Path(args.tasks_dir)
    out = results / "report"
    shutil.rmtree(out, ignore_errors=True)
    (out / "img").mkdir(parents=True)
    name = args.name or results.name
    settings = json.loads((results / "settings.json").read_text()) if (results / "settings.json").exists() else {}
    date = args.date or datetime.date.fromtimestamp((results / "results.json").stat().st_mtime).isoformat()
    label = {"a": args.a_label or settings.get("a", "A"), "b": args.b_label or settings.get("b", "B")}
    image = {"a": settings.get("a", "?"), "b": settings.get("b", "?")}
    model = {}
    for side in "ab":
        m, v = settings.get(f"{side}_model") or settings.get("model", "?"), settings.get(f"{side}_variant", settings.get("variant"))
        model[side] = m + ("" if v is None else f" ({v or 'default'} effort)")
    same_model = model["a"] == model["b"]
    question = args.question or "Does version B produce better results than version A?"
    reviewer = f" ({args.judge_label})" if args.judge_label else ""
    grades, per_task = graded(results), verdicts(results, judge)
    summaries = {t: summarize(p) for t, p in per_task.items()}
    calls = {t: call(s) for t, s in summaries.items()}
    # Tasks with a difference first, then by name.
    order = sorted(summaries, key=lambda t: (calls[t][1] == "", t))
    picked = shows(args.show)
    runs_n = settings.get("runs") or max((g.get("a", {}).get("runs", 0) for g in grades.values()), default=0)
    wins = {w: [t for t, (_, cw) in calls.items() if cw == w] for w in ("a", "b", "")}
    agree = sum(s["consistent"] for s in summaries.values()), sum(s["pairs"] for s in summaries.values())
    answer = (f"Of {len(summaries)} tasks, A was better in {len(wins['a'])}, B in {len(wins['b'])}, "
              f"and {len(wins[''])} showed no real difference.")
    prompts = {t: first_sentence((tasks_dir / t / "prompt.md").read_text()) if (tasks_dir / t / "prompt.md").exists() else ""
               for t in summaries}
    suffix = f' Each request ended with: “{settings["prompt_suffix"]}”' if settings.get("prompt_suffix") else ""

    def g_(t, side, key, scale=1):
        v = grades.get(t, {}).get(side, {}).get(key)
        return None if v is None else v * scale

    def anchor(i, t):
        """The id GitLab and GitHub give the heading of task section 3.i."""
        return re.sub(r"[^\w\- ]", "", f"3.{i} {t}: {calls[t][0]}".lower()).replace(" ", "-")

    md = [f"# {question}", "", f"*A/B test report · {name} · {date}*", "", f"**{answer}**", "",
          (f"> **In short.** Two versions of konrad, A ({label['a']}) and B ({label['b']}), did the same {len(summaries)} "
           f"tasks, {runs_n} times each, with the same AI model. " if same_model else
           f"> **In short.** Two AI models, A ({label['a']}) and B ({label['b']}), did the same {len(summaries)} tasks "
           f"in konrad, {runs_n} times each. ") + "Every result was checked automatically for hard facts, "
          "and every pair of results was compared by AI reviewers who didn't know which version made which.", "",
          "| task | outcome | pairs: A better / no difference / B better | checks passed A / B | tokens read A / B (k) |",
          "| --- | --- | --- | --- | --- |"]
    for i, t in enumerate(order, 1):
        s = summaries[t]
        md.append(f"| [{t}](#{anchor(i, t)}) | {calls[t][0]} | {s['a']} / {s['tie']} / {s['b']} | "
                  f"{fmt(g_(t, 'a', 'score', 100), '.0f')}% / {fmt(g_(t, 'b', 'score', 100), '.0f')}% | "
                  f"{fmt(g_(t, 'a', 'tokens_in', 1e-3), '.0f')} / {fmt(g_(t, 'b', 'tokens_in', 1e-3), '.0f')} |")
    md += ["", "## 1. What was compared", "", "| | A | B |", "| --- | --- | --- |",
           f"| version | {cell(label['a'])} | {cell(label['b'])} |", f"| image | `{image['a']}` | `{image['b']}` |",
           *([] if same_model else [f"| model | `{model['a']}` | `{model['b']}` |"]), "",
           ((f"Both versions ran with `{model['a']}`." if same_model else "") + f"{suffix} The tasks:").lstrip(), ""]
    md += [f"- **{t}**: “{prompts[t]}”" for t in order]
    md += ["", "## 2. How", "",
           f"1. **Same task, same input, {runs_n} runs per version.** Run 1 of A and run 1 of B form a *pair*, and so on.",
           "1. **Automatic checks.** Each task has a small test that checks fixed facts about the result (is the old name "
           "gone, is the chart there). They tell whether a result is correct, not whether it is good.",
           f"1. **Blind review.** An AI reviewer{reviewer} gets the request, the input and both results as page images and "
           "text, labelled only “X” and “Y” in random order. It scores both on *request* (did it do what was asked), "
           "*fidelity* (is the data right) and *presentation* (could you hand it on as it is), and says which it would "
           "rather receive, or that it makes no difference.",
           "1. **Each pair is reviewed twice,** by different reviewers, with X and Y swapped. A pair counts as a win only if "
           f"both reviews pick the same version, otherwise as no difference. Here they agreed in {agree[0]} of {agree[1]} pairs.",
           "1. **A task's outcome** is “clearly better” when one version wins at least half of all pairs more than the "
           "other, “somewhat better” when it wins at least two more, and “no real difference” otherwise.", "",
           "## 3. Results", "",
           "Per task: the outcome, what a reviewer saw, and where one version won, the pages that show it (A left, B right). "
           "All numbers and more reasons are folded away under each task."]

    for i, t in enumerate(order, 1):
        s, g = summaries[t], grades.get(t, {})
        verdict, winner = calls[t]
        if t in picked:
            run, pages = picked[t]
        else:
            run = next((n for n in examples(s, limit=len(s["both"])) if s["final"][n] == (winner or "tie")),
                       next(iter(examples(s, 1)), None))
            pages = list(range(1, args.pages + 1)) if winner else []
        md += ["", f"### 3.{i} {t}: {verdict}", "", f"*“{prompts[t]}”*", "",
               f"A better in {s['a']} pairs, no difference in {s['tie']}, B better in {s['b']}; the two reviews agreed "
               f"in {s['consistent']} of {s['pairs']}. Automatic checks passed: A {fmt(g_(t, 'a', 'score', 100), '.0f')}%, "
               f"B {fmt(g_(t, 'b', 'score', 100), '.0f')}%."]
        if run is not None and run in s["both"]:
            o = s["both"][run]
            md += ["", f"**A reviewer on run {run}:** {ab_terms(o[1]['reasons'], o[1]['X'])}"]
            dirs = side_dirs(judge, o)
            for doc in sorted({d.name for side in "ab" for d in dirs[side].iterdir() if d.is_dir()}):
                kind = "slide" if doc.endswith(".pptx") else "page"
                files = {side: sorted((dirs[side] / doc).glob("page-*.png")) for side in "ab"}
                for side in "ab":
                    missing = dirs[side] / doc / "MISSING.txt"
                    if missing.exists() and pages:
                        md += ["", f"*{side.upper()} has no {doc}: {missing.read_text().strip()}*"]
                for pg in pages:
                    if not any(pg <= len(files[side]) for side in "ab"):
                        continue
                    f = f"{t}-run{run}-{doc.replace('.', '-')}-{pg}.jpg"
                    pair_image(*(files[side][pg - 1] if pg <= len(files[side]) else None for side in "ab"), out / "img" / f)
                    md += ["", f"![run {run}, {doc}, {kind} {pg}: A left, B right](img/{f})", "",
                           f"*Run {run}, {doc}, {kind} {pg}: A left, B right.*"]
        def row(title, f):
            return f"| {title} | " + " | ".join(f(g[side]) if side in g else "–" for side in "ab") + " |"
        md += ["", "<details><summary>All numbers</summary>", "", "| | A | B |", "| --- | --- | --- |",
               row("automatic checks passed", lambda x: f"{x['score'] * 100:.0f}% ({x['all_pass']} of {x['runs']} runs flawless)"),
               *(f"| review score: {c} | {fmt(s['means'][(c, 'a')], '.1f')} / 5 | {fmt(s['means'][(c, 'b')], '.1f')} / 5 |"
                 for c in CRITERIA),
               row("time per run", lambda x: f"{x['seconds']:.0f} s"),
               row("tokens read / written per run", lambda x: f"{x['tokens_in'] / 1000:.0f}k / {x['tokens_out'] / 1000:.1f}k"),
               row("cost per run", lambda x: f"${x['cost']:.4f}"),
               row("runs that used a skill", lambda x: f"{x['skill_runs']} of {x['runs']}"),
               "", "</details>"]
        more = [n for n in examples(s, limit=4) if n != run][:3]
        if more:
            md += ["", "<details><summary>More reviewers' reasons</summary>", ""]
            for n_ in more:
                o = s["both"][n_]
                outcome = {"a": "A better", "b": "B better", "tie": "no difference"}[s["final"][n_]]
                md.append(f"- **Run {n_}, {outcome}:** {ab_terms(o[1]['reasons'], o[1]['X'])}")
            md += ["", "</details>"]

    md += ["", "## 4. Limits", "", f"- {runs_n} pairs per task show clear differences, not subtle ones.",
           "- The reviewers are AI models. They agreed with themselves as shown above, but nobody checked every verdict by hand.",
           (f"- The runs used `{model['a']}`; another model, such as the local one konrad is tuned for, may behave differently."
            if same_model else "- Each model ran at the effort shown, through one provider; the same model served elsewhere "
            "(another provider, a local runtime, a quantized build) may behave differently."),
           "- The reviewers are kept blind by instruction: they could have looked at other files on disk.", ""]
    (out / "README.md").write_text("\n".join(md), encoding="utf-8")
    size = sum(f.stat().st_size for f in out.rglob("*") if f.is_file())
    print(f"{answer}\n{out}/README.md, {len(list((out / 'img').iterdir()))} images, {size / 1e6:.2f} MB")


def main():
    ap = argparse.ArgumentParser()
    sub = ap.add_subparsers(dest="cmd", required=True)
    b = sub.add_parser("blind")
    b.add_argument("results")
    b.add_argument("judge")
    b.add_argument("tasks_dir")
    b.add_argument("host_judge")
    b.add_argument("--batch", type=int, default=5)
    b.add_argument("--dpi", type=int, default=60)
    b.add_argument("--tasks", default="", help="comma-separated; default all tasks in the results dir")
    r = sub.add_parser("report")
    r.add_argument("results")
    r.add_argument("judge")
    r.add_argument("--name", default="")
    r.add_argument("--tasks-dir", dest="tasks_dir", default="/ab")
    r.add_argument("--question", default="")
    r.add_argument("--a-label", dest="a_label", default="")
    r.add_argument("--b-label", dest="b_label", default="")
    r.add_argument("--judge-label", dest="judge_label", default="")
    r.add_argument("--date", default="", help="the runs' date; default: when results.json was written")
    r.add_argument("--pages", type=int, default=2, help="pages shown where one side won, unless --show picks them")
    r.add_argument("--show", action="append", default=[], metavar="TASK:RUN:PAGE[,PAGE]",
                   help="the pair and pages to show for a task, picked by hand; repeatable")
    args = ap.parse_args()
    {"blind": blind, "report": report}[args.cmd](args)


main()
