# SPDX-FileCopyrightText: 2026 Jan-Luca Bauß
# SPDX-License-Identifier: AGPL-3.0-or-later
"""Grade xlsx-edit-budget: a row and a column added to budget.xlsx (see prompt.md)."""

import re
import sys
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "lib"))
from abgrade import close, is_formula, recalc, run  # noqa: E402

from openpyxl import load_workbook  # noqa: E402

TASK = Path(__file__).resolve().parent
NAMES = [
    "budget-updated.xlsx opens", "same sheets, same order", "Utilities right below Office Rent",
    "Utilities values", "Utilities total is a correct formula", "Total row includes Utilities",
    "Total row still formulas", "Avg header right after Total", "Avg cells are formulas",
    "Avg values correct", "other rows unchanged", "number formats kept", "Notes sheet unchanged",
    "chart kept", "chart covers Utilities", "budget.xlsx untouched",
]
UTIL = [300, 300, 320, 310]


def charts(path):
    with zipfile.ZipFile(path) as z:
        return sum(1 for n in z.namelist() if n.startswith("xl/charts/chart"))


def chart_rows(path):
    """Last row of every range a chart in the workbook plots."""
    with zipfile.ZipFile(path) as z:
        xml = "".join(z.read(n).decode() for n in z.namelist() if n.startswith("xl/charts/chart"))
    refs = re.findall(r"<(?:c:)?f>([^<]*)</(?:c:)?f>", xml)
    return [int(m.group(1)) for r in refs if (m := re.search(r":\$?[A-Z]+\$?(\d+)$", r))]


def grade(ws_dir, c):
    src = TASK / "input/budget.xlsx"
    path = ws_dir / "budget-updated.xlsx"
    if not c.add(NAMES[0], path.exists(), "missing" if not path.exists() else ""):
        raise RuntimeError("no budget-updated.xlsx")
    orig, origv = load_workbook(src), load_workbook(src, data_only=True)
    wb, vb = load_workbook(path), load_workbook(recalc(path), data_only=True)
    c.add(NAMES[1], wb.sheetnames == orig.sheetnames, wb.sheetnames)
    ws, vs, ows, ovs = wb["Budget"], vb["Budget"], orig["Budget"], origv["Budget"]

    labels = {str(ws.cell(r, 1).value or "").strip(): r for r in range(1, ws.max_row + 1)}
    u, rent, total = labels.get("Utilities"), labels.get("Office Rent"), labels.get("Total")
    c.add(NAMES[2], u is not None and rent is not None and u == rent + 1, f"Office Rent {rent}, Utilities {u}")
    if u is None or total is None:
        raise RuntimeError("no Utilities or Total row")
    c.add(NAMES[3], [vs.cell(u, col).value for col in range(2, 6)] == UTIL,
          [vs.cell(u, col).value for col in range(2, 6)])
    c.add(NAMES[4], is_formula(ws.cell(u, 6).value) and close(vs.cell(u, 6).value, sum(UTIL)),
          f"{ws.cell(u, 6).value} = {vs.cell(u, 6).value}")
    want = [ovs.cell(7, col).value + add for col, add in zip(range(2, 7), [*UTIL, sum(UTIL)])]
    got = [vs.cell(total, col).value for col in range(2, 7)]
    c.add(NAMES[5], all(close(a, b) for a, b in zip(got, want)), f"{got} vs {want}")
    c.add(NAMES[6], all(is_formula(ws.cell(total, col).value) for col in range(2, 7)),
          [ws.cell(total, col).value for col in range(2, 7)])

    c.add(NAMES[7], str(ws.cell(1, 7).value or "").strip().lower() in ("avg", "average"), ws.cell(1, 7).value)
    rows = [r for r in range(2, total + 1)]
    c.add(NAMES[8], all(is_formula(ws.cell(r, 7).value) for r in rows),
          f"{sum(is_formula(ws.cell(r, 7).value) for r in rows)}/{len(rows)}")
    c.add(NAMES[9], all(close(vs.cell(r, 7).value, (vs.cell(r, 6).value or 0) / 4) for r in rows),
          [vs.cell(r, 7).value for r in rows])

    kept, fmts = True, True
    for r in range(2, 7):
        name = ows.cell(r, 1).value
        nr = labels.get(name)
        if nr is None:
            kept = False
            continue
        kept &= [vs.cell(nr, col).value for col in range(2, 7)] == [ovs.cell(r, col).value for col in range(2, 7)]
        fmts &= all(ws.cell(nr, col).number_format == ows.cell(r, col).number_format for col in range(2, 7))
    c.add(NAMES[10], kept)
    c.add(NAMES[11], fmts and ws.cell(1, 1).font.bold, ws.cell(2, 2).number_format)
    notes = [[x.value for x in row] for row in wb["Notes"].iter_rows()] if "Notes" in wb.sheetnames else None
    c.add(NAMES[12], notes == [[x.value for x in row] for row in orig["Notes"].iter_rows()])
    c.add(NAMES[13], charts(path) >= charts(src), f"{charts(path)} chart(s), input has {charts(src)}")
    ends = chart_rows(path)
    c.add(NAMES[14], ends and all(e == total - 1 for e in ends), f"ranges end at rows {ends}, last category {total - 1}")
    c.add(NAMES[15], (ws_dir / "budget.xlsx").read_bytes() == src.read_bytes())


run(grade, NAMES)
