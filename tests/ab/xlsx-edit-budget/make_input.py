# SPDX-FileCopyrightText: 2026 Jan-Luca Bauß
# SPDX-License-Identifier: AGPL-3.0-or-later
"""Regenerate input/budget.xlsx (committed, so every trial starts from the same bytes).

Run in a konrad image: openpyxl writes the workbook, then a LibreOffice round
trip stores computed formula results the way a file saved by a spreadsheet app
has them. The bar chart is the part an openpyxl re-save loses.

    podman run --rm -v "$PWD/tests/ab/xlsx-edit-budget:/t:z" --entrypoint /opt/venv/bin/python \\
        konrad:local /t/make_input.py /t/input/budget.xlsx
"""

import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill

ROWS = [
    ("Salaries", 42000, 42000, 43500, 43500),
    ("Office Rent", 6500, 6500, 6500, 6500),
    ("Software", 1800, 1750, 2100, 1900),
    ("Travel", 900, 2400, 1200, 650),
    ("Marketing", 3000, 2500, 4200, 3800),
]

wb = Workbook()
ws = wb.active
ws.title = "Budget"
ws.append(["Category", "Jan", "Feb", "Mar", "Apr", "Total"])
for i, row in enumerate(ROWS, start=2):
    ws.append([*row, f"=SUM(B{i}:E{i})"])
last = len(ROWS) + 1
ws.append(["Total", *[f"=SUM({c}2:{c}{last})" for c in "BCDEF"]])
for cell in ws[1]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="3F7A57")
for cell in ws[last + 1]:
    cell.font = Font(bold=True)
for row in ws.iter_rows(min_row=2, min_col=2):
    for cell in row:
        cell.number_format = "#,##0"
ws.column_dimensions["A"].width = 16
chart = BarChart()
chart.title = "Total per category"
chart.add_data(Reference(ws, min_col=6, min_row=1, max_row=last), titles_from_data=True)
chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=last))
ws.add_chart(chart, "I2")

notes = wb.create_sheet("Notes")
notes.append(["Budget FY2026, first four months"])
notes.append(["Bürobedarf is planned under Software from May."])
notes.append(["Owner: finance team"])

out = Path(sys.argv[1])
tmp = Path(tempfile.mkdtemp())
wb.save(tmp / "budget.xlsx")
subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx:Calc MS Excel 2007 XML",
                "--outdir", str(tmp / "out"), str(tmp / "budget.xlsx")],
               check=True, capture_output=True, env={"HOME": str(tmp), "PATH": "/usr/bin:/bin"})
shutil.copyfile(tmp / "out" / "budget.xlsx", out)
