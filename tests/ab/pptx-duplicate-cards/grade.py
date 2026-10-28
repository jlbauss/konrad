# SPDX-FileCopyrightText: 2026 Jan-Luca Bauß
# SPDX-License-Identifier: AGPL-3.0-or-later
"""Grade pptx-duplicate-cards: a profile card copied once per person, slide numbers added (see prompt.md)."""

import csv
import io
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "lib"))
from abgrade import norm, pdf_text, run, to_pdf  # noqa: E402

from PIL import Image  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER  # noqa: E402

TASK = Path(__file__).resolve().parent
NAMES = [
    "team-v2.pptx opens", "five slides", "names in the titles, in order", "roles", "role formatting kept",
    "every card has a working photo", "photo where the template had it", "card layout kept", "speaker notes",
    "template card gone", "slide numbers on every card", "no slide number on the title slide",
    "LibreOffice renders five pages", "inputs untouched",
]
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def others(slide):
    title = slide.shapes.title
    return [s for s in slide.shapes if s.has_text_frame and (title is None or s.shape_id != title.shape_id)
            and not (s.is_placeholder and s.placeholder_format.type == PP_PLACEHOLDER.SLIDE_NUMBER)]


def numbered(slide):
    return any(s.is_placeholder and s.placeholder_format.type == PP_PLACEHOLDER.SLIDE_NUMBER for s in slide.shapes) \
        or any(f.get("type") == "slidenum" for f in slide._element.iter(f"{A}fld"))


def photo_ok(pic):
    try:
        Image.open(io.BytesIO(pic.image.blob)).verify()
        return True
    except Exception:  # noqa: BLE001 — a broken relationship or bytes that aren't an image
        return False


def grade(ws_dir, c):
    src = TASK / "input/team.pptx"
    path = ws_dir / "team-v2.pptx"
    if not c.add(NAMES[0], path.exists(), "missing" if not path.exists() else ""):
        raise RuntimeError("no team-v2.pptx")
    deck, orig = Presentation(path), Presentation(src)
    people = list(csv.reader((TASK / "input/people.csv").read_text(encoding="utf-8").splitlines()))[1:]
    slides = list(deck.slides)
    cards = slides[1:]
    c.add(NAMES[1], len(slides) == 1 + len(people), len(slides))
    titles = [norm(s.shapes.title.text) if s.shapes.title is not None else None for s in slides]
    c.add(NAMES[2], titles == ["Our Team", *(norm(n) for n, _ in people)], titles)

    roles = [[norm(s.text_frame.text) for s in others(card) if s.text_frame.text.strip()] for card in cards]
    c.add(NAMES[3], len(cards) == len(people) and all(r == [norm(role)] for r, (_, role) in zip(roles, people)),
          roles)
    tpl = list(orig.slides)[1]
    tpl_run = others(tpl)[0].text_frame.paragraphs[0].runs[0].font
    fonts = [[(r.font.size, r.font.color.rgb if r.font.color and r.font.color.type else None)
              for s in others(card) for p in s.text_frame.paragraphs for r in p.runs if r.text.strip()]
             for card in cards]
    c.add(NAMES[4], fonts and all(f and all(x == (tpl_run.size, tpl_run.color.rgb) for x in f) for f in fonts),
          fonts[:1])

    pics = [[s for s in card.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE] for card in cards]
    c.add(NAMES[5], cards and all(p and all(photo_ok(x) for x in p) for p in pics),
          [[photo_ok(x) for x in p] for p in pics])
    tp = next(s for s in tpl.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE)
    c.add(NAMES[6], cards and all(p and (p[0].left, p[0].top, p[0].width) == (tp.left, tp.top, tp.width)
                                  for p in pics))
    c.add(NAMES[7], cards and all(card.slide_layout.name == tpl.slide_layout.name for card in cards),
          [card.slide_layout.name for card in cards])
    notes = [norm(card.notes_slide.notes_text_frame.text) if card.has_notes_slide else "" for card in cards]
    c.add(NAMES[8], notes == [norm(f"Intro: {n}") for n, _ in people], notes)
    c.add(NAMES[9], "Name" not in titles and not any(r == ["Role"] for r in roles))
    c.add(NAMES[10], cards and all(numbered(card) for card in cards), [numbered(card) for card in cards])
    c.add(NAMES[11], slides and not numbered(slides[0]))
    pdf = to_pdf(path)
    c.add(NAMES[12], pdf is not None and pdf_text(pdf).count("\f") == 1 + len(people), "no PDF" if pdf is None else "")
    c.add(NAMES[13], all((ws_dir / n).read_bytes() == (TASK / "input" / n).read_bytes()
                         for n in ("team.pptx", "people.csv")))


run(grade, NAMES)
