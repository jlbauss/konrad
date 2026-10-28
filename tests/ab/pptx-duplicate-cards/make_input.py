# SPDX-FileCopyrightText: 2026 Jan-Luca Bauß
# SPDX-License-Identifier: AGPL-3.0-or-later
"""Regenerate input/team.pptx (committed, so every trial starts from the same bytes).

Run in a konrad image. Slide 2 is a profile card to copy: a picture, the name in
the title placeholder, the role in a text box. python-pptx has no slide copy,
and a copied picture without its image relationship shows as a broken frame.

    podman run --rm -v "$PWD/tests/ab/pptx-duplicate-cards:/t:z" --entrypoint /opt/venv/bin/python \\
        konrad:local /t/make_input.py /t/input/team.pptx
"""

import io
import sys

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt


def photo():
    img = Image.new("RGB", (300, 300), "#e4e7eb")
    draw = ImageDraw.Draw(img)
    draw.ellipse((100, 50, 200, 150), fill="#7b8794")
    draw.ellipse((60, 170, 240, 360), fill="#7b8794")
    buf = io.BytesIO()
    img.save(buf, "PNG")
    buf.seek(0)
    return buf


def main(out):
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    title = deck.slides.add_slide(deck.slide_layouts[0])
    title.shapes.title.text = "Our Team"
    title.placeholders[1].text = "Who to call for what"

    card = deck.slides.add_slide(deck.slide_layouts[5])
    card.shapes.title.text = "Name"
    card.shapes.add_picture(photo(), Inches(1), Inches(2), width=Inches(3.5))
    role = card.shapes.add_textbox(Inches(5.5), Inches(3), Inches(7), Inches(1)).text_frame
    role.text = "Role"
    role.paragraphs[0].runs[0].font.size = Pt(28)
    role.paragraphs[0].runs[0].font.color.rgb = RGBColor(0x48, 0x65, 0x81)
    card.notes_slide.notes_text_frame.text = "Intro: Name"
    deck.save(out)


main(sys.argv[1])
